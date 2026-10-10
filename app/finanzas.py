# -*- coding: utf-8 -*-
"""Finanzas: módulo independiente. Nunca abre ni modifica cartera.json."""
import csv
import datetime as dt
import io
import json
import os
import shutil
import threading
import uuid
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from flask import Blueprint, Response, jsonify, request

TIPOS = {"ingreso", "gasto", "devolucion", "inversion", "transferencia"}
ACCIONES = {"compra", "venta"}
BASE_CATEGORIAS = {
    "gasto": {"Vivienda": ["Alquiler", "Suministros", "Mantenimiento"], "Alimentación": ["Supermercado", "Comida para llevar"], "Restaurantes y ocio": ["Restaurantes", "Bares", "Entretenimiento"], "Transporte": ["Transporte público", "Taxi", "Combustible", "Viajes"], "Salud y deporte": ["Farmacia", "Gimnasio", "Deporte"], "Compras": ["Ropa", "Electrónica", "Hogar"], "Suscripciones": ["Software", "Plataformas"], "Educación": ["Formación", "Libros"], "Impuestos y comisiones": [], "Otros gastos": []},
    "ingreso": {"Trabajo": ["Nómina", "Trabajos puntuales"], "Becas y ayudas": [], "Rendimientos financieros": ["Dividendos", "Intereses"], "Otros ingresos": []},
}
LOCK = threading.RLock()


def nuevo():
    cats = []
    for tipo, grupos in BASE_CATEGORIAS.items():
        for nombre, subs in grupos.items():
            cid = uuid.uuid4().hex
            cats.append({"id": cid, "tipo": tipo, "nombre": nombre, "subcategorias": [{"id": uuid.uuid4().hex, "nombre": n} for n in subs]})
    return {"version": 1, "movimientos": [], "categorias": cats, "cuentas": []}


def montar(app, datos):
    """Único enlace con el servidor original. Rutas y almacenamiento exclusivos."""
    bp = Blueprint("finanzas", __name__, url_prefix="/api/finanzas")
    root = os.path.join(datos, "finanzas")
    archivo = os.path.join(root, "finanzas.json")
    backups = os.path.join(root, "copias")

    def leer():
        if not os.path.isfile(archivo):
            return nuevo()
        with open(archivo, encoding="utf-8") as f:
            d = json.load(f)
        if d.get("version") != 1:
            raise ValueError("Versión de Finanzas incompatible")
        return d

    def guardar(data):
        os.makedirs(root, exist_ok=True)
        if os.path.isfile(archivo):
            os.makedirs(backups, exist_ok=True)
            nombre = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".json"
            shutil.copy2(archivo, os.path.join(backups, nombre))
            anteriores = sorted(os.listdir(backups))
            for viejo in anteriores[:-50]:
                os.remove(os.path.join(backups, viejo))
        temporal = archivo + ".tmp"
        try:
            with open(temporal, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            os.replace(temporal, archivo)
        finally:
            if os.path.exists(temporal):
                os.remove(temporal)

    def numero(x):
        try:
            n = Decimal(str(x).replace(",", "."))
            if not n.is_finite() or n <= 0:
                raise ValueError()
            return str(n.quantize(Decimal(".01"), rounding=ROUND_HALF_UP))
        except (InvalidOperation, ValueError):
            raise ValueError("El importe debe ser positivo y válido")

    def validar(x, d, previo=None):
        if not isinstance(x, dict):
            raise ValueError("Movimiento incorrecto")
        tipo = x.get("tipo")
        if tipo not in TIPOS:
            raise ValueError("Tipo desconocido")
        try:
            fecha = dt.date.fromisoformat(x["fecha"]).isoformat()
        except (KeyError, ValueError, TypeError):
            raise ValueError("Fecha incorrecta")
        concepto = str(x.get("concepto", "")).strip()[:200]
        if not concepto:
            raise ValueError("Falta concepto")
        categoria = x.get("categoria") or ""
        sub = x.get("subcategoria") or ""
        ct = "ingreso" if tipo == "ingreso" else "gasto" if tipo in ("gasto", "devolucion") else None
        if ct:
            c = next((a for a in d["categorias"] if a["id"] == categoria and a["tipo"] == ct), None)
            if not c:
                raise ValueError("Categoría inexistente")
            if sub and sub not in [a["id"] for a in c["subcategorias"]]:
                raise ValueError("Subcategoría inexistente")
        else:
            categoria, sub = "", ""
        cuenta = x.get("cuenta") or ""
        if cuenta and not any(a["id"] == cuenta for a in d["cuentas"]):
            raise ValueError("Cuenta inexistente")
        destino = x.get("destino") or ""
        if destino and not any(a["id"] == destino for a in d["cuentas"]):
            raise ValueError("Cuenta destino inexistente")
        accion = x.get("accion") or ""
        if tipo == "inversion" and accion not in ACCIONES:
            raise ValueError("La inversión necesita compra o venta")
        if tipo != "inversion":
            accion = ""
        original = x.get("original") or ""
        if tipo == "devolucion" and original:
            m = next((v for v in d["movimientos"] if v["id"] == original and v["tipo"] == "gasto"), None)
            if not m:
                raise ValueError("Gasto original inexistente")
        else:
            original = ""
        ahora = dt.datetime.now().isoformat(timespec="seconds")
        return {"id": previo["id"] if previo else uuid.uuid4().hex, "fecha": fecha, "tipo": tipo,
                "concepto": concepto, "importe": numero(x.get("importe")), "categoria": categoria,
                "subcategoria": sub, "cuenta": cuenta, "destino": destino if tipo == "transferencia" else "",
                "accion": accion, "original": original, "notas": str(x.get("notas") or "")[:1000],
                "creado": previo["creado"] if previo else ahora, "modificado": ahora}

    def totales(ms):
        out = {"ingresos": 0, "gastos": 0, "inversion": 0}
        for m in ms:
            n = int(Decimal(m["importe"]) * 100)
            if m["tipo"] == "ingreso": out["ingresos"] += n
            elif m["tipo"] == "gasto": out["gastos"] += n
            elif m["tipo"] == "devolucion": out["gastos"] -= n
            elif m["tipo"] == "inversion": out["inversion"] += n if m["accion"] == "compra" else -n
        ahorro = out["ingresos"] - out["gastos"]
        return dict(ingresos=out["ingresos"]/100, gastos=out["gastos"]/100,
                    ahorro=ahorro/100, inversion=out["inversion"]/100,
                    disponible=(ahorro-out["inversion"])/100,
                    tasa=round(100*ahorro/out["ingresos"], 2) if out["ingresos"] else None)

    def filtrar(d, filtros):
        ms = d["movimientos"]
        desde, hasta = filtros.get("desde"), filtros.get("hasta")
        if desde: ms = [m for m in ms if m["fecha"] >= desde]
        if hasta: ms = [m for m in ms if m["fecha"] <= hasta]
        for campo in ("tipo", "categoria", "subcategoria", "cuenta"):
            val = filtros.get(campo, "")
            if val: ms = [m for m in ms if m.get(campo, "") == val]
        q = filtros.get("q", "").casefold()
        if q: ms = [m for m in ms if q in m["concepto"].casefold() or q in m.get("notas", "").casefold()]
        return ms

    @bp.get("")
    def estado():
        with LOCK: return jsonify(leer())

    @bp.post("/movimientos")
    def agregar():
        try:
            with LOCK:
                d = leer()
                entradas = request.get_json()
                if not isinstance(entradas, list): entradas = [entradas]
                if not entradas or len(entradas) > 5000: raise ValueError("Lote no válido")
                nuevos = [validar(m, d) for m in entradas]
                d["movimientos"].extend(nuevos)
                guardar(d)
                return jsonify(ok=True, movimientos=nuevos)
        except (ValueError, TypeError, KeyError) as e: return jsonify(error=str(e)), 400

    @bp.put("/movimientos/<mid>")
    def editar(mid):
        try:
            with LOCK:
                d = leer()
                pos = next(i for i, m in enumerate(d["movimientos"]) if m["id"] == mid)
                d["movimientos"][pos] = validar(request.get_json(), d, d["movimientos"][pos])
                guardar(d)
                return jsonify(ok=True)
        except StopIteration: return jsonify(error="Movimiento no encontrado"), 404
        except (ValueError, TypeError, KeyError) as e: return jsonify(error=str(e)), 400

    @bp.delete("/movimientos/<mid>")
    def borrar(mid):
        with LOCK:
            d = leer()
            if any(m.get("original") == mid for m in d["movimientos"]):
                return jsonify(error="Hay devoluciones vinculadas: desvincúlalas primero"), 400
            original = len(d["movimientos"])
            d["movimientos"] = [m for m in d["movimientos"] if m["id"] != mid]
            if len(d["movimientos"]) == original: return jsonify(error="No encontrado"), 404
            guardar(d)
            return jsonify(ok=True)

    @bp.get("/analisis")
    def analisis():
        with LOCK:
            d = leer()
            ms = filtrar(d, request.args)
            grupos = defaultdict(list)
            for m in ms: grupos[m["fecha"][:7]].append(m)
            return jsonify(total=totales(ms), meses={k:totales(v) for k,v in sorted(grupos.items())},
                           categorias=d["categorias"], movimientos=ms)

    @bp.route("/configuracion/<tipo>", methods=["POST", "PUT", "DELETE"])
    def configuracion(tipo):
        if tipo not in ("cuentas", "categorias"): return jsonify(error="Tipo desconocido"), 404
        try:
            with LOCK:
                d = leer()
                val = request.get_json() or {}
                registros = d[tipo]
                rid = val.get("id", "")
                actual = next((a for a in registros if a["id"] == rid), None)
                if request.method == "DELETE":
                    if not actual: raise ValueError("Elemento inexistente")
                    if tipo == "cuentas" and any(m.get("cuenta") == rid or m.get("destino") == rid for m in d["movimientos"]):
                        raise ValueError("Cuenta en uso")
                    if tipo == "categorias" and any(m.get("categoria") == rid for m in d["movimientos"]):
                        raise ValueError("Categoría en uso")
                    registros.remove(actual)
                else:
                    nombre = str(val.get("nombre") or "").strip()[:80]
                    if not nombre: raise ValueError("Falta el nombre")
                    if request.method == "POST":
                        if tipo == "categorias":
                            if val.get("tipo") not in ("gasto", "ingreso"): raise ValueError("Tipo de categoría incorrecto")
                            registros.append({"id": uuid.uuid4().hex, "nombre": nombre, "tipo": val["tipo"], "subcategorias": []})
                        else:
                            registros.append({"id": uuid.uuid4().hex, "nombre": nombre})
                    else:
                        if not actual: raise ValueError("Elemento inexistente")
                        actual["nombre"] = nombre
                        if tipo == "categorias" and "subcategorias" in val:
                            subs = val["subcategorias"]
                            if not isinstance(subs, list): raise ValueError("Subcategorías incorrectas")
                            anteriores = {a["id"] for a in actual["subcategorias"]}
                            nuevas = [{"id": s.get("id") if s.get("id") in anteriores else uuid.uuid4().hex,
                                       "nombre": str(s.get("nombre") or "").strip()[:80]} for s in subs]
                            if any(not s["nombre"] for s in nuevas): raise ValueError("Nombre de subcategoría vacío")
                            usadas = {m.get("subcategoria") for m in d["movimientos"] if m.get("categoria")==rid}
                            if not usadas.issubset({s["id"] for s in nuevas} | {""}):
                                raise ValueError("Subcategoría en uso")
                            actual["subcategorias"] = nuevas
                guardar(d)
                return jsonify(ok=True, registros=registros)
        except (ValueError, TypeError, KeyError) as e: return jsonify(error=str(e)), 400

    @bp.get("/copias")
    def lista_copias():
        with LOCK:
            return jsonify(copias=sorted(os.listdir(backups), reverse=True) if os.path.isdir(backups) else [])

    @bp.post("/recuperar")
    def recuperar():
        nombre = (request.get_json() or {}).get("nombre", "")
        if not isinstance(nombre, str) or "/" in nombre or "\\" in nombre or not nombre.endswith(".json"):
            return jsonify(error="Nombre incorrecto"), 400
        with LOCK:
            origen = os.path.join(backups, nombre)
            if not os.path.isfile(origen): return jsonify(error="Copia inexistente"), 404
            with open(origen, encoding="utf-8") as f: d = json.load(f)
            if d.get("version") != 1: return jsonify(error="Versión incompatible"), 400
            guardar(d)
            return jsonify(ok=True)

    @bp.get("/copia.json")
    def descarga():
        with LOCK: data = leer()
        return Response(json.dumps(data, ensure_ascii=False, indent=2), mimetype="application/json",
                        headers={"Content-Disposition": 'attachment; filename="finanzas.json"'})

    @bp.post("/restaurar-json")
    def restaurar_json():
        try:
            d = request.get_json()
            if not isinstance(d, dict) or d.get("version") != 1 or not all(isinstance(d.get(k), list) for k in ("movimientos", "categorias", "cuentas")):
                raise ValueError("Copia no válida")
            with LOCK:
                guardar(d)
            return jsonify(ok=True)
        except (ValueError, TypeError) as e: return jsonify(error=str(e)), 400

    @bp.get("/exportar")
    def exportar():
        with LOCK: d = leer()
        ms = filtrar(d, request.args)
        columnas = [c for c in request.args.get("campos", "fecha,tipo,concepto,importe,categoria,subcategoria,cuenta,accion,notas").split(",") if c in {"fecha","tipo","concepto","importe","categoria","subcategoria","cuenta","accion","notas"}]
        if not columnas: return jsonify(error="Sin columnas"), 400
        formatos = request.args.get("formato", "csv")
        nombres = {a["id"]:a["nombre"] for a in d["categorias"]+d["cuentas"]}
        nombres.update({s["id"]:s["nombre"] for a in d["categorias"] for s in a["subcategorias"]})
        filas = [[nombres.get(m.get(c),m.get(c,"")) if c in ("categoria","subcategoria","cuenta") else m.get(c,"") for c in columnas] for m in ms]
        if formatos == "xlsx":
            from openpyxl import Workbook
            from flask import send_file
            wb=Workbook();ws=wb.active;ws.title="Movimientos";ws.append(columnas)
            for row in filas: ws.append(row)
            b=io.BytesIO();wb.save(b);b.seek(0)
            return send_file(b,as_attachment=True,download_name="finanzas.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if formatos != "csv": return jsonify(error="Formato no válido"), 400
        b=io.StringIO();w=csv.writer(b,delimiter=";");w.writerow(columnas);w.writerows(filas)
        return Response("\ufeff"+b.getvalue(),mimetype="text/csv",headers={"Content-Disposition":'attachment; filename="finanzas.csv"'})

    @bp.post("/importar-vista-previa")
    def vista_previa():
        f=request.files.get("archivo")
        if not f: return jsonify(error="Falta archivo"), 400
        try:
            if f.filename.lower().endswith(".xlsx"):
                from openpyxl import load_workbook
                wb=load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True)
                filas=list(wb.active.values)
            elif f.filename.lower().endswith(".csv"):
                texto=f.read().decode("utf-8-sig")
                dialecto=csv.Sniffer().sniff(texto[:4096], delimiters=";,|\t")
                filas=list(csv.reader(io.StringIO(texto), dialecto))
            else: raise ValueError("Usa CSV o XLSX")
            if not filas: raise ValueError("Archivo vacío")
            return jsonify(columnas=[str(v or "") for v in filas[0]],filas=[[str(v) if v is not None else "" for v in fila] for fila in filas[1:5001]])
        except Exception as e: return jsonify(error=str(e)), 400

    app.register_blueprint(bp)
